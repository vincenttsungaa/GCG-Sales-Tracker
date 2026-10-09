from io import BytesIO

from openpyxl import load_workbook

from lib.export import build_workbook


def test_user_text_is_never_a_formula():
    data = build_workbook([{"kind": "card", "name": '=HYPERLINK("x")', "status": "sold"}])
    ws = load_workbook(BytesIO(data))["Listings"]
    assert ws["A2"].value == '=HYPERLINK("x")' and ws["A2"].data_type == "s"
    assert ws["N2"].data_type == "f"  # the app's own totals stay formulas


def test_every_text_column_is_guarded():
    data = build_workbook([{"kind": "item", "name": "x", "part": "=1+1", "edition": "=2+2", "card_quantities": {"=A1": 1}}])
    ws = load_workbook(BytesIO(data))["Listings"]
    formulas = {c.column_letter for c in ws[2] if c.data_type == "f"}
    assert formulas == {"N", "S", "T"}  # Asking total, Sale total, Profit only

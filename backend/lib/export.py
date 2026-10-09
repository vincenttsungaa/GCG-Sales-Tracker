"""Listings → an editable Excel workbook (.xlsx), and sending it by email.

The workbook is plain .xlsx (Office Open XML), so it opens and edits in Microsoft Excel,
Google Sheets, LibreOffice / OpenOffice, Apple Numbers and the Excel / Sheets phone apps.
Values are real numbers and dates (not text), totals and the Summary sheet are formulas,
so editing a price or a quantity updates them.
"""

from __future__ import annotations

import os
import smtplib
from datetime import datetime
from email.message import EmailMessage
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
AUD = '"A$"#,##0.00'
STATUS_LABELS = {"for_sale": "For Sale", "pending": "Pending", "on_hold": "Storage", "sold": "Sold"}

# (header, width) — order of the Listings sheet
COLUMNS = [
    ("Name", 42),
    ("Notes", 36),
    ("Type", 9),
    ("Code", 12),
    ("Category / card type", 20),
    ("Rarity", 8),
    ("Colour", 9),
    ("Edition", 16),
    ("Part", 22),
    ("Contents", 36),
    ("Status", 11),
    ("Qty", 6),
    ("Asking price (each)", 16),
    ("Asking total", 14),
    ("Purchase price (each)", 18),
    ("Buyer", 18),
    ("Deal date", 12),
    ("Sale price (each)", 15),
    ("Sale total", 13),
    ("Profit", 12),
    ("Condition", 14),
    ("Added", 12),
]
COL = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(COLUMNS)}
OWN_FORMULAS = {COL["Asking total"], COL["Sale total"], COL["Profit"]}


def _title(value: str | None) -> str:
    return (value or "").replace("_", " ").title()


def _contents(item: dict) -> str:
    """What's inside a listing: the products in a bundle, or the cards / designs picked."""
    bundle = item.get("bundle_items") or []
    if len(bundle) > 1:
        parts = []
        for e in bundle:
            label = e.get("code") or e.get("name") or "Item"
            qty = e.get("quantity") or 1
            parts.append(f"{label} ({qty}x)" + (f" — {e['detail']}" if e.get("detail") else ""))
        return "; ".join(parts)
    copies = item.get("card_quantities") or {}
    if copies:
        return ", ".join(f"{k} ({v}x)" for k, v in copies.items())
    picked = [*(item.get("resource_cards") or []), *(item.get("alt_art_cards") or []), *(item.get("sleeve_designs") or [])]
    return ", ".join(picked)


def _date(value) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, str) and value:
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return None
    return None


def build_workbook(items: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Listings"
    ws.append([name for name, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row, item in enumerate(items, start=2):
        is_card = item.get("kind") == "card"
        ws.append(
            [
                item.get("name"),
                item.get("notes"),
                "Card" if is_card else "Item",
                item.get("card_no") or item.get("set_code"),
                _title(item.get("card_type") if is_card else item.get("category")),
                item.get("rarity"),
                _title(item.get("color")),
                item.get("edition"),
                item.get("part"),
                _contents(item),
                STATUS_LABELS.get(item.get("status"), _title(item.get("status"))),
                item.get("quantity") or 1,
                item.get("price") or 0,
                f"={COL['Qty']}{row}*{COL['Asking price (each)']}{row}",
                item.get("purchase_price"),
                item.get("buyer_name"),
                _date(item.get("deal_date")),
                item.get("sale_price"),
                f'=IF({COL["Sale price (each)"]}{row}="","",{COL["Qty"]}{row}*{COL["Sale price (each)"]}{row})',
                # profit only for sold rows that have both a sale and a purchase price
                f'=IF(AND({COL["Status"]}{row}="Sold",{COL["Sale price (each)"]}{row}<>"",{COL["Purchase price (each)"]}{row}<>""),'
                f'{COL["Qty"]}{row}*({COL["Sale price (each)"]}{row}-{COL["Purchase price (each)"]}{row}),"")',
                item.get("condition"),
                _date(item.get("created_at")),
            ]
        )
        for cell in ws[row]:  # typed text starting with "=" stays text; only the app's totals are formulas
            if cell.column_letter not in OWN_FORMULAS and isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    last = max(len(items) + 1, 2)
    money = ["Asking price (each)", "Asking total", "Purchase price (each)", "Sale price (each)", "Sale total", "Profit"]
    for name in money:
        for cell in ws[f"{COL[name]}2:{COL[name]}{last}"]:
            cell[0].number_format = AUD
    for name in ("Deal date", "Added"):
        for cell in ws[f"{COL[name]}2:{COL[name]}{last}"]:
            cell[0].number_format = "yyyy-mm-dd"
    for name in ("Name", "Contents", "Notes"):
        for cell in ws[f"{COL[name]}2:{COL[name]}{last}"]:
            cell[0].alignment = Alignment(wrap_text=True, vertical="top")

    # a real Excel table: header filters / sorting, banded rows, works in every spreadsheet app
    if items:
        table = Table(displayName="Listings", ref=f"A1:{get_column_letter(len(COLUMNS))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"

    # Summary — formulas over the Listings sheet, so they follow any edits
    s = wb.create_sheet("Summary")
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 16
    s.append(["CardStakk export", datetime.now().strftime("%Y-%m-%d %H:%M")])
    s["A1"].font = Font(bold=True, size=13)
    s.append([])
    s.append(["Listings", "Count"])
    rng = lambda name: f"Listings!{COL[name]}2:{COL[name]}{last}"  # noqa: E731
    for label in ("For Sale", "Pending", "Storage", "Sold"):
        s.append([label, f'=COUNTIF({rng("Status")},"{label}")'])
    s.append(["Total", "=SUM(B4:B7)"])
    s.append([])
    s.append(["Value (AUD)", "Amount"])
    s.append(["Asking value (not sold)", f'=SUMIF({rng("Status")},"<>Sold",{rng("Asking total")})'])
    s.append(["Realized sales", f'=SUMIF({rng("Status")},"Sold",{rng("Sale total")})'])
    s.append(["Profit after costs", f"=SUM({rng('Profit')})"])
    for r in (3, 10):
        for c in ("A", "B"):
            s[f"{c}{r}"].font = Font(bold=True, color="FFFFFF")
            s[f"{c}{r}"].fill = PatternFill("solid", fgColor="5C6A8A")
    for r in (11, 12, 13):
        s[f"B{r}"].number_format = AUD
    s["A8"].font = Font(bold=True)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def export_filename() -> str:
    return f"cardstakk-listings-{datetime.now().strftime('%Y-%m-%d')}.xlsx"


class EmailNotConfigured(RuntimeError):
    pass


def send_workbook(to: str, data: bytes, filename: str, count: int, note: str | None = None) -> None:
    """Email the workbook through the SMTP server set in backend/.env (SMTP_HOST, SMTP_PORT,
    SMTP_USER, SMTP_PASSWORD, SMTP_FROM; SMTP_SSL=true for port 465)."""
    host = os.environ.get("SMTP_HOST")
    if not host:
        raise EmailNotConfigured("Email isn't set up yet: add the SMTP settings to backend/.env.")
    port = int(os.environ.get("SMTP_PORT") or "587")
    user = os.environ.get("SMTP_USER")
    password = os.environ.get("SMTP_PASSWORD")
    sender = os.environ.get("SMTP_FROM") or user or "cardstakk@localhost"

    msg = EmailMessage()
    msg["Subject"] = f"CardStakk listings — {count} {'listing' if count == 1 else 'listings'}"
    msg["From"] = sender
    msg["To"] = to
    body = [
        f"Attached: {filename} ({count} {'listing' if count == 1 else 'listings'}).",
        "It opens in Excel, Google Sheets, LibreOffice, Numbers and the Excel / Sheets phone apps.",
    ]
    if note:
        body = [note.strip(), "", *body]
    msg.set_content("\n".join(body) + "\n\n— Sent from CardStakk")
    msg.add_attachment(data, maintype="application", subtype=XLSX_MIME.split("/", 1)[1], filename=filename)

    if os.environ.get("SMTP_SSL", "").lower() in ("1", "true", "yes") or port == 465:
        with smtplib.SMTP_SSL(host, port, timeout=30) as smtp:
            if user:
                smtp.login(user, password or "")
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=30) as smtp:
            smtp.starttls()
            if user:
                smtp.login(user, password or "")
            smtp.send_message(msg)
